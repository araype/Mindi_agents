"""Convierte revision_clinica.md en revision_clinica.docx para comentar en Word.

El .docx se abre con el control de cambios activado, para que las ediciones
del equipo médico queden registradas. Entiende solo el Markdown que genera
generar.py: títulos, párrafos, listas (con un nivel de anidación), citas,
tablas, **negrita**, *cursiva* y `código`.

Correr desde la raíz del repo, después de generar.py:
    python revision_clinica/generar_docx.py

Requiere python-docx (pip install "python-docx==1.1.2").
"""

import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

CARPETA = Path(__file__).parent
ENTRADA = CARPETA / "revision_clinica.md"
SALIDA = CARPETA / "revision_clinica.docx"

GRIS_CITA = RGBColor(0x40, 0x40, 0x40)
BORDE_CITA = "A0A0A0"
FONDO_ENCABEZADO = "E8EEF4"

INLINE = re.compile(r"(\*\*[^*]+\*\*|\*[^*\s][^*]*\*|`[^`]+`)")
VINETA = re.compile(r"^(\s*)- (.*)$")
NUMERO = re.compile(r"^(\d+)\. (.*)$")


def _texto_con_formato(parrafo, texto: str) -> None:
    for trozo in INLINE.split(texto):
        if not trozo:
            continue
        if trozo.startswith("**") and trozo.endswith("**"):
            parrafo.add_run(trozo[2:-2]).bold = True
        elif trozo.startswith("`") and trozo.endswith("`"):
            run = parrafo.add_run(trozo[1:-1])
            run.font.name = "Consolas"
        elif trozo.startswith("*") and trozo.endswith("*") and len(trozo) > 2:
            parrafo.add_run(trozo[1:-1]).italic = True
        else:
            parrafo.add_run(trozo)


# Elementos que el esquema exige DESPUÉS de pBdr dentro de pPr.
_DESPUES_DE_PBDR = [
    "shd", "tabs", "suppressAutoHyphens", "kinsoku", "wordWrap",
    "overflowPunct", "topLinePunct", "autoSpaceDE", "autoSpaceDN", "bidi",
    "adjustRightInd", "snapToGrid", "spacing", "ind", "contextualSpacing",
    "mirrorIndents", "suppressOverlap", "jc", "textDirection", "textAlignment",
    "textboxTightWrap", "outlineLvl", "divId", "cnfStyle", "rPr", "sectPr",
    "pPrChange",
]


def _borde_izquierdo(parrafo) -> None:
    propiedades = parrafo._p.get_or_add_pPr()
    bordes = OxmlElement("w:pBdr")
    izquierdo = OxmlElement("w:left")
    for clave, valor in (("w:val", "single"), ("w:sz", "12"),
                         ("w:space", "8"), ("w:color", BORDE_CITA)):
        izquierdo.set(qn(clave), valor)
    bordes.append(izquierdo)
    siguientes = {qn(f"w:{nombre}") for nombre in _DESPUES_DE_PBDR}
    for hijo in propiedades:
        if hijo.tag in siguientes:
            hijo.addprevious(bordes)
            return
    propiedades.append(bordes)


def _sombrear_celda(celda, color: str) -> None:
    propiedades = celda._tc.get_or_add_tcPr()
    sombra = OxmlElement("w:shd")
    for clave, valor in (("w:val", "clear"), ("w:color", "auto"), ("w:fill", color)):
        sombra.set(qn(clave), valor)
    propiedades.append(sombra)


def _activar_control_de_cambios(documento) -> None:
    """Word abre el archivo registrando los cambios (Revisar > Control de cambios)."""
    ajustes = documento.settings.element
    # La plantilla de python-docx trae <w:zoom> sin el atributo obligatorio.
    zoom = ajustes.find(qn("w:zoom"))
    if zoom is not None and zoom.get(qn("w:percent")) is None:
        zoom.set(qn("w:percent"), "100")
    if ajustes.find(qn("w:trackRevisions")) is not None:
        return
    marca = OxmlElement("w:trackRevisions")
    # El esquema exige que vaya antes de defaultTabStop.
    tabulacion = ajustes.find(qn("w:defaultTabStop"))
    if tabulacion is not None:
        tabulacion.addprevious(marca)
    else:
        ajustes.append(marca)


def _numero_de_pagina(documento) -> None:
    pie = documento.sections[0].footer.paragraphs[0]
    pie.alignment = 1  # centrado
    run = pie.add_run()
    for tipo, texto in (("begin", None), (None, "PAGE"), ("end", None)):
        if tipo:
            campo = OxmlElement("w:fldChar")
            campo.set(qn("w:fldCharType"), tipo)
        else:
            campo = OxmlElement("w:instrText")
            campo.set(qn("xml:space"), "preserve")
            campo.text = texto
        run._r.append(campo)


def _preparar(documento) -> None:
    seccion = documento.sections[0]
    seccion.page_width, seccion.page_height = Cm(21), Cm(29.7)  # A4
    for margen in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(seccion, margen, Cm(2.5))
    normal = documento.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    documento.core_properties.title = "Revisión clínica de Mindi"
    documento.core_properties.author = "Equipo Mindi"
    _activar_control_de_cambios(documento)
    _numero_de_pagina(documento)


def _agregar_tabla(documento, filas: list[list[str]]) -> None:
    columnas = len(filas[0])
    tabla = documento.add_table(rows=len(filas), cols=columnas)
    tabla.style = "Table Grid"
    tabla.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, fila in enumerate(filas):
        for j, texto in enumerate(fila):
            celda = tabla.cell(i, j)
            parrafo = celda.paragraphs[0]
            parrafo.paragraph_format.space_after = Pt(0)
            _texto_con_formato(parrafo, texto)
            if i == 0:
                for run in parrafo.runs:
                    run.bold = True
                _sombrear_celda(celda, FONDO_ENCABEZADO)
    documento.add_paragraph()


def convertir(markdown: str):
    documento = Document()
    _preparar(documento)

    parrafo = []      # líneas del párrafo en curso
    cita = []         # párrafos de la cita en curso
    tabla = []        # filas de la tabla en curso
    item = None       # [estilo, texto, sangría] del ítem de lista en curso

    def cerrar_parrafo():
        if parrafo:
            _texto_con_formato(documento.add_paragraph(), " ".join(parrafo))
            parrafo.clear()

    def cerrar_item():
        nonlocal item
        if item:
            estilo, texto, sangria = item
            p = documento.add_paragraph(style=estilo)
            if sangria:
                p.paragraph_format.left_indent = Cm(1.27 + 0.63 * sangria)
            if estilo == "Normal":  # ítem numerado: número literal y sangría francesa
                p.paragraph_format.left_indent = Cm(0.9)
                p.paragraph_format.first_line_indent = Cm(-0.6)
            _texto_con_formato(p, texto)
            item = None

    def cerrar_cita():
        bloques, actual = [], []
        for linea in cita:
            if linea:
                actual.append(linea)
            elif actual:
                bloques.append(actual)
                actual = []
        if actual:
            bloques.append(actual)
        for bloque in bloques:
            p = documento.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.8)
            _borde_izquierdo(p)
            _texto_con_formato(p, " ".join(bloque))
            for run in p.runs:
                run.font.color.rgb = GRIS_CITA
        cita.clear()

    def cerrar_tabla():
        if tabla:
            _agregar_tabla(documento, tabla)
            tabla.clear()

    def cerrar_todo():
        cerrar_parrafo()
        cerrar_item()
        cerrar_cita()
        cerrar_tabla()

    for linea in markdown.splitlines():
        if linea.startswith(">"):
            cerrar_parrafo(); cerrar_item(); cerrar_tabla()
            cita.append(linea[1:].strip())
            continue
        if cita:
            cerrar_cita()

        if linea.startswith("|"):
            cerrar_parrafo(); cerrar_item()
            celdas = [c.strip() for c in linea.strip().strip("|").split("|")]
            if not all(set(c) <= {"-", ":"} for c in celdas):
                tabla.append(celdas)
            continue
        cerrar_tabla()

        titulo = re.match(r"^(#+) (.*)$", linea)
        if titulo:
            cerrar_todo()
            nivel = len(titulo.group(1))
            if nivel == 1:
                documento.add_heading(titulo.group(2), level=0)
            else:
                documento.add_heading(titulo.group(2), level=min(nivel - 1, 6))
            continue

        vineta = VINETA.match(linea)
        numero = NUMERO.match(linea)
        if vineta:
            cerrar_parrafo(); cerrar_item()
            sangria = len(vineta.group(1)) // 2
            estilo = "List Bullet 2" if sangria else "List Bullet"
            item = [estilo, vineta.group(2), 0]
            continue
        if numero:
            cerrar_parrafo(); cerrar_item()
            item = ["Normal", f"{numero.group(1)}.\t{numero.group(2)}", 0]
            continue

        if not linea.strip():
            cerrar_parrafo(); cerrar_item()
            continue

        if item and linea.startswith(" "):  # continuación del ítem
            item[1] += " " + linea.strip()
            continue
        cerrar_item()
        parrafo.append(linea.strip())

    cerrar_todo()
    return documento


if __name__ == "__main__":
    convertir(ENTRADA.read_text(encoding="utf-8")).save(SALIDA)
    print(f"-> {SALIDA}")
